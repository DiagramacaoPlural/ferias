from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, abort, flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user

from ..alertas import prazos_dos_funcionarios
from ..consultas import (
    ciclo_atual,
    ciclo_por_ano,
    ciclos_disponiveis,
    ferias_no_intervalo,
    funcionarios_ativos,
    grade_cronograma,
    minimo_no_setor,
    painel_cobertura,
    serie_cobertura,
)
from ..models import (
    STATUS_AGUARDANDO_FUNCIONARIO,
    STATUS_AGUARDANDO_GERENTE,
    STATUS_AGUARDANDO_LIDER,
    STATUS_AJUSTE_GERENTE,
    STATUS_AJUSTE_LIDER,
    STATUS_APROVADA,
    STATUS_PROGRAMADA,
    Ferias,
    Funcionario,
    SugestaoFerias,
    db,
)
from ..workflow import (
    aceitar_proposta,
    enviar_ao_lider,
    gerente_aprovar,
    gerente_sugerir,
    lider_definir,
    lider_sugerir,
    lider_validar,
    montar_calendario,
    recusar_proposta,
)
from ..regras import (
    AVISO_13_FEVEREIRO,
    DIAS_MAX_CICLO,
    analisar_periodo_db,
    calcular_fim,
    meses_do_ciclo,
    nivel_cobertura,
    pode_adiantar_13,
    rotulo_ciclo,
    sobrepoe,
    ultimo_dia_mes,
)
from . import registrar, requer

bp = Blueprint("ferias", __name__, url_prefix="/ferias")


def _ciclo_da_requisicao() -> tuple[date, date]:
    return ciclo_por_ano(request.args.get("ciclo", type=int))


def _pendencias_do_funcionario(funcionario_id: int) -> list[Ferias]:
    return list(
        db.session.scalars(
            db.select(Ferias)
            .where(
                Ferias.funcionario_id == funcionario_id,
                Ferias.status.in_(
                    (STATUS_AGUARDANDO_FUNCIONARIO, STATUS_AJUSTE_GERENTE, STATUS_AJUSTE_LIDER)
                ),
            )
            .order_by(Ferias.inicio)
        )
    )


def _pode_editar(ferias: Ferias) -> bool:
    return current_user.admin or not ferias.aprovada


def _ler_formulario():
    erros = []
    funcionario = db.session.get(Funcionario, request.form.get("funcionario_id", type=int) or 0)
    if not funcionario:
        erros.append("Selecione o funcionário.")
    try:
        inicio = datetime.strptime(request.form.get("inicio", ""), "%Y-%m-%d").date()
    except ValueError:
        inicio = None
        erros.append("Informe uma data de início válida.")
    dias = request.form.get("dias", type=int)
    if not dias:
        erros.append("Informe a quantidade de dias.")
    return funcionario, inicio, dias, erros


def _quer_adiantar_13(fonte) -> bool:
    return str(fonte.get("adianta_13") or "").strip().lower() in ("1", "true", "on")


def _quer_vende_10(fonte) -> bool:
    return str(fonte.get("vende_10") or "").strip().lower() in ("1", "true", "on")


def _destino_calendario():
    return request.referrer or url_for("ferias.calendario")


def _erro_adianta_13(inicio: date | None, adianta: bool) -> str | None:
    if adianta and inicio and not pode_adiantar_13(inicio):
        return AVISO_13_FEVEREIRO
    return None


@bp.route("/")
@requer("gestao")
def lista():
    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    funcionario_id = request.args.get("funcionario", type=int)
    status = request.args.get("status", "")

    todas = [f for f in ferias_no_intervalo(ciclo_ini, ciclo_fim) if ciclo_ini <= f.inicio <= ciclo_fim]
    ferias = todas
    if funcionario_id:
        ferias = [f for f in ferias if f.funcionario_id == funcionario_id]
    if status in (STATUS_APROVADA, STATUS_PROGRAMADA):
        ferias = [f for f in ferias if f.status == status]

    avisos, alternativas = {}, {}
    for f in ferias:
        analise = analisar_periodo_db(f.funcionario, f.inicio, f.dias, f.id)
        avisos[f.id] = analise.avisos
        if analise.alternativas:
            alternativas[f.id] = analise.alternativas

    ativos = funcionarios_ativos()
    min_disp = minimo_no_setor()
    cobertura = serie_cobertura(todas, ciclo_ini, ciclo_fim, len(ativos), min_disp)
    painel = painel_cobertura(cobertura, len(ativos), min_disp)

    return render_template(
        "ferias/lista.html",
        ferias=ferias,
        avisos=avisos,
        alternativas=alternativas,
        cobertura=cobertura,
        painel=painel,
        total_ativos=len(ativos),
        min_disponiveis=min_disp,
        ciclo_ini=ciclo_ini,
        ciclos=ciclos_disponiveis(ciclo_ini),
        rotulo_ciclo=rotulo_ciclo(ciclo_ini),
        funcionarios=db.session.scalars(db.select(Funcionario).order_by(Funcionario.nome)).all(),
        filtro_funcionario=funcionario_id,
        filtro_status=status,
    )


@bp.route("/nova", methods=["GET", "POST"])
@bp.route("/<int:ferias_id>/editar", methods=["GET", "POST"])
@requer("gestao")
def formulario(ferias_id: int | None = None):
    ferias = db.session.get(Ferias, ferias_id) if ferias_id else None
    if ferias_id and not ferias:
        abort(404)
    if ferias and not _pode_editar(ferias):
        flash("Férias aprovadas só podem ser alteradas por um administrador.", "aviso")
        return redirect(url_for("ferias.lista"))

    avisos, erros, alternativas = [], [], []
    if request.method == "POST":
        funcionario, inicio, dias, erros = _ler_formulario()
        adianta_13 = _quer_adiantar_13(request.form)
        vende_10 = _quer_vende_10(request.form)
        if erro_13 := _erro_adianta_13(inicio, adianta_13):
            erros.append(erro_13)
        if not erros:
            analise = analisar_periodo_db(funcionario, inicio, dias, ferias_id)
            erros, avisos, alternativas = analise.erros, analise.avisos, analise.alternativas

        confirmado = request.form.get("confirmar") == "1"
        if not erros and (not avisos or confirmado):
            novo = ferias is None
            if novo:
                ferias = Ferias(criado_por=current_user.login)
                db.session.add(ferias)
            ferias.funcionario = funcionario
            ferias.definir_periodo(inicio, dias)
            ferias.observacao = request.form.get("observacao", "").strip() or None
            ferias.adianta_13 = adianta_13
            ferias.vende_10 = vende_10

            if current_user.admin:
                aprovar = request.form.get("aprovada") == "1"
                if aprovar and not ferias.aprovada:
                    ferias.aprovado_por, ferias.aprovado_em = current_user.login, datetime.now()
                ferias.status = STATUS_APROVADA if aprovar else STATUS_PROGRAMADA
                if not aprovar:
                    ferias.aprovado_por = ferias.aprovado_em = None

            descricao = f"{funcionario.nome}: {inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y} ({dias} dias)"
            if adianta_13:
                descricao += " · 1ª parcela do 13º"
            if vende_10:
                descricao += " · vende 10 dias"
            if avisos:
                descricao += f" | salvo com {len(avisos)} aviso(s)"
            registrar("Férias criadas" if novo else "Férias alteradas", descricao)
            db.session.commit()
            flash("Férias salvas com sucesso.", "sucesso")
            return redirect(url_for("ferias.lista", ciclo=ciclo_atual(inicio)[0].year))

    funcionarios = funcionarios_ativos()
    if ferias and ferias.funcionario not in funcionarios:
        funcionarios.append(ferias.funcionario)

    return render_template(
        "ferias/formulario.html",
        ferias=ferias,
        funcionarios=funcionarios,
        erros=erros,
        avisos=avisos,
        alternativas=alternativas,
        dados=request.form if request.method == "POST" else None,
        dias_max=DIAS_MAX_CICLO,
    )


@bp.route("/funcionario/<int:funcionario_id>/dados")
@requer("calendario")
def dados_funcionario(funcionario_id: int):
    """Dados do ciclo para preencher o formulário ao trocar o funcionário."""
    funcionario = db.session.get(Funcionario, funcionario_id)
    if not funcionario:
        abort(404)
    referencia = None
    bruto = request.args.get("inicio", "")
    if bruto:
        try:
            referencia = datetime.strptime(bruto, "%Y-%m-%d").date()
        except ValueError:
            referencia = None
    ciclo_ini, ciclo_fim = ciclo_atual(referencia)
    if referencia is None:
        do_atual = [f for f in funcionario.ferias if ciclo_ini <= f.inicio <= ciclo_fim]
        if not do_atual:
            futuras = [f for f in funcionario.ferias if f.inicio > ciclo_fim]
            if futuras:
                ciclo_ini, ciclo_fim = ciclo_atual(futuras[0].inicio)
    _, prox_fim = ciclo_atual(ciclo_fim + timedelta(days=1))
    periodos = [f for f in funcionario.ferias if ciclo_ini <= f.inicio <= prox_fim]
    usados = sum(f.dias for f in funcionario.ferias if ciclo_ini <= f.inicio <= ciclo_fim)
    prazo = prazos_dos_funcionarios([funcionario]).get(funcionario.id)
    return jsonify(
        nome=funcionario.nome,
        admissao=funcionario.data_admissao.strftime("%d/%m/%Y") if funcionario.data_admissao else None,
        limite=prazo.limite.strftime("%d/%m/%Y") if prazo and prazo.limite else None,
        usados=usados,
        restante=max(DIAS_MAX_CICLO - usados, 0),
        periodos=[
            {
                "id": f.id,
                "inicio": f.inicio.isoformat(),
                "fim": f.fim.strftime("%d/%m/%Y"),
                "dias": f.dias,
                "observacao": f.observacao or "",
                "adianta_13": bool(f.adianta_13),
                "vende_10": bool(f.vende_10),
                "aprovada": f.aprovada,
                "status": f.status_rotulo,
                "no_ciclo": ciclo_ini <= f.inicio <= ciclo_fim,
            }
            for f in periodos
        ],
    )


@bp.route("/analisar", methods=["POST"])
@requer("calendario")
def analisar():
    """Validação ao vivo usada pelo formulário enquanto o usuário digita."""
    dados = request.get_json(silent=True) or {}
    funcionario = db.session.get(Funcionario, int(dados.get("funcionario_id") or 0))
    try:
        inicio = datetime.strptime(dados.get("inicio", ""), "%Y-%m-%d").date()
        dias = int(dados.get("dias"))
    except (TypeError, ValueError):
        return jsonify(fim=None, erros=[], avisos=[])

    fim = calcular_fim(inicio, dias) if dias > 0 else None
    erros_13 = []
    if erro_13 := _erro_adianta_13(inicio, _quer_adiantar_13(dados)):
        erros_13.append(erro_13)
    if not funcionario:
        return jsonify(fim=fim.strftime("%d/%m/%Y") if fim else None, erros=erros_13, avisos=[])

    ferias_id = int(dados["ferias_id"]) if str(dados.get("ferias_id") or "").isdigit() else None
    analise = analisar_periodo_db(funcionario, inicio, dias, ferias_id)
    return jsonify(
        fim=fim.strftime("%d/%m/%Y") if fim else None,
        erros=erros_13 + analise.erros,
        avisos=analise.avisos,
        alternativas=[
            {"inicio": a.inicio.isoformat(), "fim": a.fim.strftime("%d/%m/%Y"), "rotulo": a.inicio.strftime("%d/%m/%Y")}
            for a in analise.alternativas
        ],
    )


@bp.route("/<int:ferias_id>/aprovar", methods=["POST"])
@requer("admin")
def alternar_aprovacao(ferias_id: int):
    ferias = db.get_or_404(Ferias, ferias_id)
    if ferias.aprovada:
        ferias.status, ferias.aprovado_por, ferias.aprovado_em = STATUS_PROGRAMADA, None, None
        acao = "Aprovação removida"
    else:
        ferias.status, ferias.aprovado_por, ferias.aprovado_em = STATUS_APROVADA, current_user.login, datetime.now()
        acao = "Férias aprovadas"
    registrar(acao, f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y}")
    db.session.commit()
    flash(f"{acao}: {ferias.funcionario.nome}.", "sucesso")
    return redirect(request.referrer or url_for("ferias.lista"))


@bp.route("/lote", methods=["POST"])
@requer("gestao")
def acao_lote():
    ids = request.form.getlist("ids", type=int)
    acao = request.form.get("acao", "")
    destino = request.referrer or url_for("ferias.lista")
    if not ids:
        flash("Selecione pelo menos um período.", "aviso")
        return redirect(destino)

    registros = db.session.scalars(db.select(Ferias).where(Ferias.id.in_(ids))).all()
    if acao == "aprovar":
        if not current_user.admin:
            abort(403)
        feitos = []
        for ferias in registros:
            if ferias.aprovada:
                continue
            ferias.status = STATUS_APROVADA
            ferias.aprovado_por = current_user.login
            ferias.aprovado_em = datetime.now()
            feitos.append(ferias.funcionario.nome)
        registrar("Férias aprovadas em lote", ", ".join(feitos)[:500] if feitos else "nenhum")
        db.session.commit()
        flash(f"{len(feitos)} período(s) aprovado(s)." if feitos else "Os períodos selecionados já estavam aprovados.", "sucesso" if feitos else "info")
        return redirect(destino)

    if acao in ("13_sim", "13_nao", "vende_sim", "vende_nao"):
        sim = acao in ("13_sim", "vende_sim")
        vende = acao.startswith("vende_")
        feitos, recusados = [], []
        for ferias in registros:
            if not _pode_editar(ferias):
                recusados.append(ferias.funcionario.nome)
                continue
            if not vende and sim and _erro_adianta_13(ferias.inicio, True):
                recusados.append(ferias.funcionario.nome)
                continue
            if vende:
                ferias.vende_10 = sim
            else:
                ferias.adianta_13 = sim
            feitos.append(ferias.funcionario.nome)
        rotulo = "Venda de 10 dias" if vende else "13º"
        registrar(
            f"{rotulo} atualizado em lote",
            ("Sim" if sim else "Não") + (": " + ", ".join(feitos) if feitos else ""),
        )
        db.session.commit()
        if feitos:
            flash(f"{rotulo} marcado como {'Sim' if sim else 'Não'} em {len(feitos)} período(s).", "sucesso")
        if recusados:
            extra = ". A 1ª parcela só vale para férias a partir de fevereiro." if sim and not vende else "."
            flash("Não foi possível alterar: " + ", ".join(recusados) + extra, "aviso")
        if not feitos and not recusados:
            flash("Nenhum período foi alterado.", "aviso")
        return redirect(destino)

    flash("Ação inválida.", "erro")
    return redirect(destino)


@bp.route("/<int:ferias_id>/mover", methods=["POST"])
@requer("gestao")
def mover(ferias_id: int):
    ferias = db.get_or_404(Ferias, ferias_id)
    if not _pode_editar(ferias):
        flash("Férias aprovadas só podem ser alteradas por um administrador.", "aviso")
        return redirect(url_for("ferias.lista"))
    try:
        inicio = datetime.strptime(request.form.get("inicio", ""), "%Y-%m-%d").date()
    except ValueError:
        flash("Data alternativa inválida.", "erro")
        return redirect(url_for("ferias.lista"))

    analise = analisar_periodo_db(ferias.funcionario, inicio, ferias.dias, ferias.id)
    if analise.erros:
        flash(analise.erros[0], "erro")
        return redirect(url_for("ferias.lista"))

    anterior = f"{ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y}"
    ferias.definir_periodo(inicio, ferias.dias)
    registrar(
        "Férias reagendadas",
        f"{ferias.funcionario.nome}: {anterior} → {ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y}",
    )
    db.session.commit()
    flash(f"Férias de {ferias.funcionario.nome} movidas para {ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y}.", "sucesso")
    return redirect(request.referrer or url_for("ferias.lista"))


@bp.route("/<int:ferias_id>/excluir", methods=["POST"])
@requer("gestao")
def excluir(ferias_id: int):
    ferias = db.get_or_404(Ferias, ferias_id)
    if not _pode_editar(ferias):
        flash("Férias aprovadas só podem ser excluídas por um administrador.", "aviso")
        return redirect(url_for("ferias.lista"))
    registrar(
        "Férias excluídas",
        f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y} ({ferias.dias} dias)",
    )
    db.session.delete(ferias)
    db.session.commit()
    flash("Férias excluídas.", "sucesso")
    return redirect(request.referrer or url_for("ferias.lista"))


@bp.route("/calendario")
@requer("calendario")
def calendario():
    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    ferias = ferias_no_intervalo(ciclo_ini, ciclo_fim, apenas_ativos=True)

    meses = []
    for primeiro in meses_do_ciclo(ciclo_ini):
        ultimo = ultimo_dia_mes(primeiro)
        do_mes = [f for f in ferias if sobrepoe(f.inicio, f.fim, primeiro, ultimo)]
        if not do_mes and primeiro.month != 11:
            continue
        sai_no_mes = [f for f in do_mes if primeiro <= f.inicio <= ultimo]
        meses.append(
            {
                "data": primeiro,
                "aprovadas": [f for f in sai_no_mes if f.aprovada],
                "programadas": [f for f in sai_no_mes if not f.aprovada],
                "atual": primeiro.year == date.today().year and primeiro.month == date.today().month,
                "de_ferias": 0,
                "disp": 0,
                "nivel_disp": "pouco",
            }
        )

    grade = {}
    saidas = [m for m in meses if m["aprovadas"] or m["programadas"]]
    if saidas:
        grade = grade_cronograma(saidas[0]["data"], ultimo_dia_mes(meses[-1]["data"]))
        total = grade.get("total_ativos") or 0
        min_setor = grade.get("min_disponiveis") or 0
        ausentes = grade.get("ausentes") or {}
        dias_grade = grade.get("dias") or []
        for mes in meses:
            dias_mes = [d for d in dias_grade if d.year == mes["data"].year and d.month == mes["data"].month]
            if dias_mes and total:
                pior = max(len(ausentes.get(d, [])) for d in dias_mes)
                mes["de_ferias"] = pior
                mes["disp"] = total - pior
                mes["nivel_disp"] = nivel_cobertura(pior, total, min_setor)
            else:
                mes["de_ferias"] = 0
                mes["disp"] = total
                mes["nivel_disp"] = "pouco"

    minhas = _pendencias_do_funcionario(current_user.funcionario_id) if current_user.funcionario_id else []
    fila_lider = []
    if current_user.pode_lider():
        fila_lider = [
            f
            for f in db.session.scalars(
                db.select(Ferias).where(Ferias.status == STATUS_AGUARDANDO_LIDER).order_by(Ferias.inicio)
            )
            if f.funcionario and f.funcionario.ativo
        ]
    fila_gerente = []
    if current_user.pode_gerente():
        fila_gerente = [
            f
            for f in db.session.scalars(
                db.select(Ferias).where(Ferias.status == STATUS_AGUARDANDO_GERENTE).order_by(Ferias.inicio)
            )
            if f.funcionario and f.funcionario.ativo
        ]

    return render_template(
        "ferias/calendario.html",
        meses=meses,
        ciclo_ini=ciclo_ini,
        ciclos=ciclos_disponiveis(ciclo_ini),
        rotulo_ciclo=rotulo_ciclo(ciclo_ini),
        minhas_pendencias=minhas,
        fila_lider=fila_lider,
        fila_gerente=fila_gerente,
        **grade,
    )


def _sugestoes_do_ciclo(ano: int) -> dict[int, list[SugestaoFerias]]:
    por_func: dict[int, list[SugestaoFerias]] = {}
    for s in db.session.scalars(
        db.select(SugestaoFerias).where(SugestaoFerias.ciclo_ano == ano).order_by(
            SugestaoFerias.funcionario_id, SugestaoFerias.ordem
        )
    ):
        por_func.setdefault(s.funcionario_id, []).append(s)
    return por_func


@bp.route("/opcoes", methods=["GET", "POST"])
@requer("calendario")
def opcoes():
    funcionario = current_user.funcionario
    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    if request.method == "POST" and request.form.get("ciclo", type=int):
        ciclo_ini, ciclo_fim = ciclo_por_ano(request.form.get("ciclo", type=int))
    existentes = []
    pendencias = []
    if funcionario:
        existentes = db.session.scalars(
            db.select(SugestaoFerias)
            .where(SugestaoFerias.funcionario_id == funcionario.id, SugestaoFerias.ciclo_ano == ciclo_ini.year)
            .order_by(SugestaoFerias.ordem)
        ).all()
        pendencias = _pendencias_do_funcionario(funcionario.id)

    avisos, erros, alternativas = [], [], []
    if request.method == "POST":
        if not funcionario:
            flash("Seu usuário não está ligado a um funcionário cadastrado.", "erro")
            return redirect(url_for("ferias.opcoes"))
        lidas = []
        for ordem in (1, 2):
            bruto = request.form.get(f"inicio_{ordem}", "").strip()
            if not bruto:
                continue
            try:
                inicio = datetime.strptime(bruto, "%Y-%m-%d").date()
            except ValueError:
                erros.append(f"Opção {ordem}: data de início inválida.")
                continue
            dias = request.form.get(f"dias_{ordem}", type=int)
            if not dias:
                erros.append(f"Opção {ordem}: informe a quantidade de dias.")
                continue
            adianta_13 = _quer_adiantar_13({"adianta_13": request.form.get(f"adianta_13_{ordem}")})
            if erro_13 := _erro_adianta_13(inicio, adianta_13):
                erros.append(f"Opção {ordem}: {erro_13}")
                continue
            analise = analisar_periodo_db(funcionario, inicio, dias)
            if analise.erros:
                erros.extend(f"Opção {ordem}: {e}" for e in analise.erros)
            avisos.extend(f"Opção {ordem}: {a}" for a in analise.avisos)
            if analise.alternativas and ordem == 1:
                alternativas = analise.alternativas
            lidas.append(
                {
                    "ordem": ordem,
                    "inicio": inicio,
                    "dias": dias,
                    "ciclo_ano": ciclo_atual(inicio)[0].year,
                    "adianta_13": adianta_13,
                    "vende_10": _quer_vende_10({"vende_10": request.form.get(f"vende_10_{ordem}")}),
                }
            )
        if not lidas and not erros:
            erros.append("Informe pelo menos uma opção de datas.")
        if not erros:
            anos = {item["ciclo_ano"] for item in lidas} | {ciclo_ini.year}
            antigas = db.session.scalars(
                db.select(SugestaoFerias).where(
                    SugestaoFerias.funcionario_id == funcionario.id,
                    SugestaoFerias.ciclo_ano.in_(anos),
                )
            ).all()
            for antigo in antigas:
                db.session.delete(antigo)
            db.session.flush()
            for item in lidas:
                s = SugestaoFerias(
                    funcionario=funcionario,
                    ciclo_ano=item["ciclo_ano"],
                    ordem=item["ordem"],
                    adianta_13=item["adianta_13"],
                    vende_10=item["vende_10"],
                )
                s.definir_periodo(item["inicio"], item["dias"])
                db.session.add(s)
            registrar("Opções de férias enviadas", f"{funcionario.nome}: {len(lidas)} opção(ões)")
            db.session.commit()
            anos_enviados = {item["ciclo_ano"] for item in lidas}
            gravadas = db.session.scalars(
                db.select(SugestaoFerias)
                .where(
                    SugestaoFerias.funcionario_id == funcionario.id,
                    SugestaoFerias.ciclo_ano.in_(anos_enviados),
                )
                .order_by(SugestaoFerias.ciclo_ano, SugestaoFerias.ordem)
            ).all()
            por_ano: dict[int, list] = {}
            for sugestao in gravadas:
                por_ano.setdefault(sugestao.ciclo_ano, []).append(sugestao)
            for opcoes_ano in por_ano.values():
                enviar_ao_lider(funcionario, opcoes_ano)
            flash("Suas opções de datas foram enviadas ao líder para validar ou sugerir outra data.", "sucesso")
            destino_ano = lidas[0]["ciclo_ano"]
            if current_user.pode_lider():
                return redirect(url_for("ferias.validacao_lider", ciclo=destino_ano))
            return redirect(url_for("ferias.opcoes", ciclo=destino_ano))

    return render_template(
        "ferias/opcoes.html",
        funcionario=funcionario,
        ciclo_ini=ciclo_ini,
        ciclos=ciclos_disponiveis(ciclo_ini),
        rotulo_ciclo=rotulo_ciclo(ciclo_ini),
        existentes=existentes,
        pendencias=pendencias,
        erros=erros,
        avisos=avisos,
        alternativas=alternativas,
        dados=request.form if request.method == "POST" else None,
        dias_max=DIAS_MAX_CICLO,
    )


@bp.route("/montar", methods=["POST"])
@requer("gestao")
def montar():
    if not (current_user.pode_lider() or current_user.admin):
        abort(403)
    ano = request.form.get("ciclo", type=int)
    resumo = montar_calendario(ano)
    if resumo["montados"]:
        flash(f"Fila do líder atualizada: {len(resumo['montados'])} período(s) para validar.", "sucesso")
    else:
        flash("Ninguém novo para a fila do líder. Quem já enviou opções já está no fluxo.", "aviso")
    return redirect(url_for("ferias.validacao_lider", ciclo=ano or ciclo_atual()[0].year))


@bp.route("/<int:ferias_id>/responder", methods=["POST"])
@requer("calendario")
def responder(ferias_id: int):
    ferias = db.get_or_404(Ferias, ferias_id)
    if ferias.funcionario_id != current_user.funcionario_id and not current_user.admin:
        abort(403)
    if ferias.status not in (STATUS_AGUARDANDO_FUNCIONARIO, STATUS_AJUSTE_GERENTE, STATUS_AJUSTE_LIDER):
        flash("Este período não está aguardando a sua resposta.", "aviso")
        return redirect(_destino_calendario())
    if request.form.get("acao") == "aceitar":
        aceitar_proposta(ferias)
        registrar("Funcionário aceitou a data", f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y}")
        flash("Data aceita. O líder vai validar e enviar ao gerente.", "sucesso")
    else:
        recusar_proposta(ferias)
        registrar("Funcionário recusou a data", f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y}")
        flash("Você recusou. O líder define a data e envia ao gerente.", "aviso")
    ano = ciclo_atual(ferias.inicio)[0].year
    if current_user.pode_lider():
        return redirect(url_for("ferias.validacao_lider", ciclo=ano))
    return redirect(url_for("ferias.calendario", ciclo=ano))


@bp.route("/<int:ferias_id>/lider", methods=["POST"])
@requer("gestao")
def acao_lider(ferias_id: int):
    if not current_user.pode_lider():
        abort(403)
    ferias = db.get_or_404(Ferias, ferias_id)
    if ferias.status != STATUS_AGUARDANDO_LIDER:
        flash("Este período não está aguardando o líder.", "aviso")
        return redirect(_destino_calendario())
    bruto = request.form.get("inicio", "").strip()
    acao = request.form.get("acao", "")
    destino_lider = url_for("ferias.validacao_lider", ciclo=ciclo_atual(ferias.inicio)[0].year)
    inicio = ferias.inicio
    if bruto:
        try:
            inicio = datetime.strptime(bruto, "%Y-%m-%d").date()
        except ValueError:
            flash("Data inválida.", "erro")
            return redirect(destino_lider)
    definir_final = bool(ferias.recusou_ajuste) or acao == "definir"
    sugerir = (not definir_final) and (acao == "sugerir" or (acao != "validar" and inicio != ferias.inicio))
    if sugerir:
        if inicio == ferias.inicio:
            flash("Informe uma data diferente para sugerir ao funcionário.", "aviso")
            return redirect(destino_lider)
        lider_sugerir(ferias, inicio, current_user.login)
        registrar("Líder sugeriu outra data", f"{ferias.funcionario.nome}: {inicio:%d/%m/%Y}")
        flash(f"Nova data enviada a {ferias.funcionario.nome} por e-mail. Ele precisa aceitar ou recusar.", "sucesso")
        return redirect(destino_lider)
    if inicio != ferias.inicio:
        lider_definir(ferias, inicio, current_user.login)
        registrar("Líder definiu a data", f"{ferias.funcionario.nome}: {inicio:%d/%m/%Y}")
    else:
        lider_validar(ferias, current_user.login)
        registrar("Líder validou", f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y}")
    if definir_final:
        flash("Data definida e encaminhada ao gerente (Aceite gerente).", "sucesso")
    else:
        flash("Validado e encaminhado ao gerente (Aceite gerente).", "sucesso")
    ano = ciclo_atual(ferias.inicio)[0].year
    if current_user.pode_gerente():
        return redirect(url_for("ferias.validacao_gerente", ciclo=ano))
    return redirect(url_for("ferias.validacao_lider", ciclo=ano))


@bp.route("/<int:ferias_id>/gerente", methods=["POST"])
@requer("gestao")
def acao_gerente(ferias_id: int):
    if not current_user.pode_gerente():
        abort(403)
    ferias = db.get_or_404(Ferias, ferias_id)
    if ferias.status != STATUS_AGUARDANDO_GERENTE:
        flash("Este período não está aguardando o gerente.", "aviso")
        return redirect(_destino_calendario())
    if request.form.get("acao") == "sugerir":
        try:
            inicio = datetime.strptime(request.form.get("inicio", ""), "%Y-%m-%d").date()
        except ValueError:
            flash("Informe a nova data sugerida.", "erro")
            return redirect(_destino_calendario())
        analise = analisar_periodo_db(ferias.funcionario, inicio, ferias.dias, ferias.id)
        if analise.erros:
            flash(analise.erros[0], "erro")
            return redirect(_destino_calendario())
        gerente_sugerir(ferias, inicio, current_user.login)
        registrar("Gerente sugeriu mudança", f"{ferias.funcionario.nome}: {inicio:%d/%m/%Y}")
        flash("Sugestão enviada ao funcionário por e-mail.", "sucesso")
    else:
        gerente_aprovar(ferias, current_user.login)
        registrar("Gerente aprovou", f"{ferias.funcionario.nome}: {ferias.inicio:%d/%m/%Y} a {ferias.fim:%d/%m/%Y}")
        flash(f"Férias de {ferias.funcionario.nome} aprovadas. E-mail enviado.", "sucesso")
    return redirect(url_for("ferias.validacao_gerente", ciclo=ciclo_atual(ferias.inicio)[0].year))


@bp.route("/validacao/lider")
@requer("gestao")
def validacao_lider():
    if not current_user.pode_lider():
        abort(403)
    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    por_func = _sugestoes_do_ciclo(ciclo_ini.year)
    ativos = funcionarios_ativos()
    enviadas = [{"funcionario": f, "opcoes": por_func[f.id]} for f in ativos if f.id in por_func]
    faltando = [f for f in ativos if f.id not in por_func]
    fila = [
        f
        for f in db.session.scalars(
            db.select(Ferias).where(Ferias.status == STATUS_AGUARDANDO_LIDER).order_by(Ferias.inicio)
        )
        if f.funcionario and f.funcionario.ativo
    ]
    aguardando_funcionario = [
        f
        for f in db.session.scalars(
            db.select(Ferias).where(Ferias.status == STATUS_AJUSTE_LIDER).order_by(Ferias.inicio)
        )
        if f.funcionario and f.funcionario.ativo
    ]
    return render_template(
        "ferias/validacao_lider.html",
        ciclo_ini=ciclo_ini,
        ciclos=ciclos_disponiveis(ciclo_ini),
        rotulo_ciclo=rotulo_ciclo(ciclo_ini),
        enviadas=enviadas,
        faltando=faltando,
        fila=fila,
        aguardando_funcionario=aguardando_funcionario,
    )


@bp.route("/validacao/gerente")
@requer("gestao")
def validacao_gerente():
    if not current_user.pode_gerente():
        abort(403)
    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    fila = [
        f
        for f in db.session.scalars(
            db.select(Ferias).where(Ferias.status == STATUS_AGUARDANDO_GERENTE).order_by(Ferias.inicio)
        )
        if f.funcionario and f.funcionario.ativo
    ]
    return render_template(
        "ferias/validacao_gerente.html",
        ciclo_ini=ciclo_ini,
        ciclos=ciclos_disponiveis(ciclo_ini),
        rotulo_ciclo=rotulo_ciclo(ciclo_ini),
        fila=fila,
    )


@bp.route("/exportar")
@requer("gestao")
def exportar():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ciclo_ini, ciclo_fim = _ciclo_da_requisicao()
    ferias = [f for f in ferias_no_intervalo(ciclo_ini, ciclo_fim) if ciclo_ini <= f.inicio <= ciclo_fim]

    wb = Workbook()
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="1F4E79")
    verde = PatternFill("solid", fgColor="E2F0D9")

    def formatar_cabecalho(ws, larguras):
        for col, largura in enumerate(larguras, start=1):
            celula = ws.cell(row=1, column=col)
            celula.font, celula.fill = cabecalho_fonte, cabecalho_fundo
            celula.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = largura
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Férias"
    ws.append(["Funcionário", "Início", "Fim", "Dias", "Status", "1ª parcela 13º", "Vende 10 dias", "Observação", "Aprovado por", "Aprovado em"])
    for f in ferias:
        ws.append([
            f.funcionario.nome, f.inicio, f.fim, f.dias, f.status_rotulo,
            "Sim" if f.adianta_13 else "Não", "Sim" if f.vende_10 else "Não",
            f.observacao or "", f.aprovado_por or "", f.aprovado_em,
        ])
        linha = ws.max_row
        ws.cell(row=linha, column=2).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=3).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=10).number_format = "DD/MM/YYYY HH:MM"
        if f.aprovada:
            for col in range(1, 11):
                ws.cell(row=linha, column=col).fill = verde
    formatar_cabecalho(ws, [24, 13, 13, 8, 13, 16, 14, 40, 16, 18])

    resumo = wb.create_sheet("Resumo por funcionário")
    resumo.append(["Funcionário", "Períodos", "Dias aprovados", "Dias programados", "Saldo (de 30)"])
    for func in db.session.scalars(db.select(Funcionario).order_by(Funcionario.nome)):
        do_func = [f for f in ferias if f.funcionario_id == func.id]
        if not func.ativo and not do_func:
            continue
        aprovados = sum(f.dias for f in do_func if f.aprovada)
        programados = sum(f.dias for f in do_func if not f.aprovada)
        resumo.append([func.nome, len(do_func), aprovados, programados, DIAS_MAX_CICLO - aprovados - programados])
    formatar_cabecalho(resumo, [24, 10, 16, 18, 14])

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    nome = f"ferias_{ciclo_ini:%Y-%m}_a_{ciclo_fim:%Y-%m}.xlsx"
    return send_file(
        arquivo,
        as_attachment=True,
        download_name=nome,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
