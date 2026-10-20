from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from sqlalchemy.exc import IntegrityError

from ..consultas import ferias_no_intervalo, funcionarios_ativos
from ..models import (
    LOCAIS_PROVA,
    TIPOS_ETAPA,
    TIPOS_LANCAMENTO,
    EtapaProva,
    Funcionario,
    LancamentoDiario,
    ProjetoProva,
    db,
)
from ..regras import DIAS_SEMANA_ABREV, MESES, sobrepoe as periodos_sobrepoem
from ..verbas import calcular_verba, taxas as carregar_taxas
from . import registrar, requer

bp = Blueprint("provas", __name__, url_prefix="/provas")

OPCOES_SEMANAS = [4, 8, 13, 26]
CORES_ETAPA = {
    "Diagramação": {"DF": "df", "SP": "sp", "": "df"},
    "Produção": {"": "producao"},
    "Manuseio FGV": {"": "manuseio"},
    "Aplicação": {"": "aplicacao"},
}


def _projetos(apenas_ativos: bool = True) -> list[ProjetoProva]:
    consulta = db.select(ProjetoProva).order_by(ProjetoProva.ordem, ProjetoProva.nome)
    if apenas_ativos:
        consulta = consulta.filter_by(ativo=True)
    return list(db.session.scalars(consulta))


def _etapas(apenas_ativos: bool = True) -> list[EtapaProva]:
    consulta = db.select(EtapaProva).join(ProjetoProva).order_by(ProjetoProva.ordem, EtapaProva.ordem, EtapaProva.id)
    if apenas_ativos:
        consulta = consulta.where(ProjetoProva.ativo.is_(True))
    return list(db.session.scalars(consulta))


def _segunda_feira(dia: date) -> date:
    return dia - timedelta(days=dia.weekday())


def _data_form(nome: str) -> date | None:
    try:
        return datetime.strptime(request.form.get(nome, ""), "%Y-%m-%d").date()
    except ValueError:
        return None


def _decimal_form(nome: str) -> Decimal | None:
    bruto = (request.form.get(nome) or "").strip().replace(",", ".")
    if not bruto:
        return None
    try:
        return Decimal(bruto)
    except InvalidOperation:
        raise ValueError("Informe um valor numérico válido para a verba.")


def _equipe_do_formulario() -> list[Funcionario]:
    ids = request.form.getlist("equipe")
    if not ids:
        return []
    return list(db.session.scalars(db.select(Funcionario).where(Funcionario.id.in_([int(i) for i in ids if i.isdigit()]))))


def _funcionarios_do_diario() -> list[Funcionario]:
    ids = [i for i in request.form.getlist("funcionarios") if str(i).isdigit()]
    if not ids:
        unico = request.form.get("funcionario_id")
        if unico and str(unico).isdigit():
            ids = [unico]
    if not ids:
        return []
    return list(
        db.session.scalars(
            db.select(Funcionario).where(Funcionario.id.in_([int(i) for i in ids])).order_by(Funcionario.nome)
        )
    )


def _dias_do_periodo(inicio: date, fim: date, so_uteis: bool) -> list[date]:
    dias = []
    dia = inicio
    while dia <= fim:
        if not so_uteis or dia.weekday() < 5:
            dias.append(dia)
        dia += timedelta(days=1)
    return dias


def conflitos_mesmo_local(etapas: list[EtapaProva]) -> list[tuple[EtapaProva, EtapaProva]]:
    datadas = [e for e in etapas if e.inicio and e.fim]
    pares = []
    for i, a in enumerate(datadas):
        for b in datadas[i + 1 :]:
            mesmo_tipo = a.tipo == b.tipo
            mesmo_local = (a.local or "") == (b.local or "")
            if mesmo_tipo and mesmo_local and periodos_sobrepoem(a.inicio, a.fim, b.inicio, b.fim):
                pares.append((a, b))
    return pares


def avisos_ferias_na_etapa(etapa: EtapaProva) -> list[str]:
    if not etapa.inicio or not etapa.fim or not etapa.equipe:
        return []
    ids = {f.id for f in etapa.equipe}
    avisos = []
    for periodo in ferias_no_intervalo(etapa.inicio, etapa.fim, apenas_ativos=True):
        if periodo.funcionario_id in ids:
            avisos.append(
                f"{periodo.funcionario.nome} está de férias de {periodo.inicio.strftime('%d/%m/%Y')} "
                f"a {periodo.fim.strftime('%d/%m/%Y')}"
            )
    return avisos


@bp.route("/")
@requer("provas")
def painel():
    hoje = date.today()
    projetos = _projetos()
    etapas = [e for p in projetos for e in p.etapas]
    andamento = [e for e in etapas if e.status_em(hoje) == "andamento"]
    futuras = sorted(
        [e for e in etapas if e.status_em(hoje) == "futura"],
        key=lambda e: e.inicio or date.max,
    )
    sem_data = [e for e in etapas if e.status_em(hoje) == "sem_data"]
    conflitos = conflitos_mesmo_local(etapas)
    avisos = []
    for etapa in andamento + futuras[:8]:
        for aviso in avisos_ferias_na_etapa(etapa):
            avisos.append(f"{etapa.rotulo_completo}: {aviso}")
    return render_template(
        "provas/painel.html",
        projetos=projetos,
        etapas=etapas,
        andamento=andamento,
        futuras=futuras[:8],
        sem_data=sem_data,
        conflitos=conflitos,
        avisos=avisos,
    )


@bp.route("/projetos")
@requer("provas")
def projetos():
    return render_template(
        "provas/projetos.html",
        projetos=_projetos(apenas_ativos=False),
        conflitos=conflitos_mesmo_local(_etapas(apenas_ativos=False)),
    )


@bp.route("/projetos/novo", methods=["GET", "POST"])
@bp.route("/projetos/<int:projeto_id>", methods=["GET", "POST"])
@requer("provas")
def projeto_form(projeto_id: int | None = None):
    projeto = db.session.get(ProjetoProva, projeto_id) if projeto_id else None
    if projeto_id and not projeto:
        abort(404)

    erros = []
    if request.method == "POST":
        nome = (request.form.get("nome") or "").strip()
        if not nome:
            erros.append("Informe o nome do projeto.")
        else:
            outro = db.session.scalar(
                db.select(ProjetoProva).filter(db.func.lower(ProjetoProva.nome) == nome.lower())
            )
            if outro and outro.id != (projeto.id if projeto else None):
                erros.append("Já existe um projeto com este nome.")
        if not erros:
            if not projeto:
                ordem = (db.session.scalar(db.select(db.func.max(ProjetoProva.ordem))) or 0) + 1
                projeto = ProjetoProva(ordem=ordem)
                db.session.add(projeto)
            projeto.nome = nome
            projeto.ativo = request.form.get("ativo") == "1"
            registrar("Projeto de prova", f"{'Novo' if not projeto_id else 'Editar'}: {nome}")
            db.session.commit()
            flash("Projeto salvo.", "sucesso")
            return redirect(url_for("provas.projeto_form", projeto_id=projeto.id))

    return render_template(
        "provas/projeto_form.html",
        projeto=projeto,
        funcionarios=funcionarios_ativos(),
        erros=erros,
        dados=request.form if erros else None,
    )


@bp.route("/projetos/<int:projeto_id>/excluir", methods=["POST"])
@requer("provas")
def excluir_projeto(projeto_id: int):
    projeto = db.session.get(ProjetoProva, projeto_id) or abort(404)
    registrar("Projeto de prova", f"Excluir: {projeto.nome}")
    db.session.delete(projeto)
    db.session.commit()
    flash("Projeto excluído.", "sucesso")
    return redirect(url_for("provas.projetos"))


@bp.route("/etapas/nova", methods=["GET", "POST"])
@bp.route("/etapas/<int:etapa_id>", methods=["GET", "POST"])
@requer("provas")
def etapa_form(etapa_id: int | None = None):
    etapa = db.session.get(EtapaProva, etapa_id) if etapa_id else None
    if etapa_id and not etapa:
        abort(404)

    projetos = _projetos(apenas_ativos=False)
    funcionarios = list(db.session.scalars(db.select(Funcionario).order_by(Funcionario.nome)))
    erros, avisos = [], []

    if request.method == "POST":
        projeto = db.session.get(ProjetoProva, request.form.get("projeto_id", type=int) or 0)
        tipo = (request.form.get("tipo") or "").strip()
        local = (request.form.get("local") or "").strip() or None
        inicio = _data_form("inicio")
        fim = _data_form("fim")
        if not projeto:
            erros.append("Selecione o projeto.")
        if tipo not in TIPOS_ETAPA:
            erros.append("Selecione o tipo da etapa.")
        if tipo == "Diagramação" and local not in LOCAIS_PROVA:
            erros.append("Informe o local (DF ou SP) para diagramação.")
        if tipo != "Diagramação":
            local = None
        if (inicio and not fim) or (fim and not inicio):
            erros.append("Informe início e término, ou deixe as duas datas em branco.")
        if inicio and fim and fim < inicio:
            erros.append("O término não pode ser anterior ao início.")
        try:
            pcs = int(request.form.get("pcs_adicionais") or 0)
            if pcs < 0:
                raise ValueError
        except ValueError:
            pcs = 0
            erros.append("PCs adicionais deve ser um número inteiro.")

        if not erros:
            if not etapa:
                ordem = len(projeto.etapas) + 1
                etapa = EtapaProva(projeto=projeto, ordem=ordem)
                db.session.add(etapa)
            etapa.projeto = projeto
            etapa.tipo = tipo
            etapa.local = local
            etapa.inicio = inicio
            etapa.fim = fim
            etapa.pcs_adicionais = pcs
            etapa.observacao = (request.form.get("observacao") or "").strip() or None
            etapa.equipe = _equipe_do_formulario()
            avisos = avisos_ferias_na_etapa(etapa)
            if avisos and not request.form.get("confirmar"):
                erros.append("Há pessoas da equipe de férias neste período. Confirme para salvar mesmo assim.")
            else:
                registrar("Etapa de prova", etapa.rotulo_completo)
                db.session.commit()
                flash("Etapa salva.", "sucesso")
                return redirect(url_for("provas.projeto_form", projeto_id=projeto.id))

    return render_template(
        "provas/etapa_form.html",
        etapa=etapa,
        projetos=projetos,
        funcionarios=funcionarios,
        tipos=TIPOS_ETAPA,
        locais=LOCAIS_PROVA,
        erros=erros,
        avisos=avisos,
        dados=request.form if request.method == "POST" else None,
        projeto_padrao=request.args.get("projeto", type=int),
    )


@bp.route("/etapas/<int:etapa_id>/excluir", methods=["POST"])
@requer("provas")
def excluir_etapa(etapa_id: int):
    etapa = db.session.get(EtapaProva, etapa_id) or abort(404)
    projeto_id = etapa.projeto_id
    registrar("Etapa de prova", f"Excluir: {etapa.rotulo_completo}")
    db.session.delete(etapa)
    db.session.commit()
    flash("Etapa excluída.", "sucesso")
    return redirect(url_for("provas.projeto_form", projeto_id=projeto_id))


@bp.route("/cronograma")
@requer("provas")
def cronograma():
    etapas = [e for e in _etapas() if e.inicio and e.fim]
    try:
        inicio = datetime.strptime(request.args.get("inicio", ""), "%Y-%m-%d").date()
    except ValueError:
        inicio = min((e.inicio for e in etapas), default=date.today())
    inicio = _segunda_feira(inicio)

    semanas = request.args.get("semanas", type=int) or 16
    semanas = min(max(semanas, 1), 52)
    fim = inicio + timedelta(days=semanas * 7 - 1)
    hoje = date.today()
    dias = [inicio + timedelta(days=i) for i in range((fim - inicio).days + 1)]

    meses_cabecalho = []
    for dia in dias:
        rotulo = f"{MESES[dia.month - 1]} {dia.year}"
        if meses_cabecalho and meses_cabecalho[-1]["rotulo"] == rotulo:
            meses_cabecalho[-1]["dias"] += 1
        else:
            meses_cabecalho.append({"rotulo": rotulo, "dias": 1})

    visiveis = [e for e in etapas if e.inicio <= fim and e.fim >= inicio]
    projetos_ids = []
    linhas = []
    for etapa in visiveis:
        if etapa.projeto_id not in projetos_ids:
            projetos_ids.append(etapa.projeto_id)
            linhas.append({"tipo": "projeto", "projeto": etapa.projeto, "celulas": []})
        celulas = []
        for dia in dias:
            ativa = etapa.inicio <= dia <= etapa.fim
            celulas.append({"dia": dia, "ativa": ativa, "etapa": etapa if ativa else None})
        classe = CORES_ETAPA.get(etapa.tipo, {}).get(etapa.local or "", "df")
        linhas.append({"tipo": "etapa", "etapa": etapa, "celulas": celulas, "classe": classe})

    df_dia = {dia: [] for dia in dias}
    sp_dia = {dia: [] for dia in dias}
    for etapa in visiveis:
        if etapa.tipo != "Diagramação":
            continue
        alvo = df_dia if etapa.local == "DF" else sp_dia if etapa.local == "SP" else None
        if not alvo:
            continue
        for dia in dias:
            if etapa.inicio <= dia <= etapa.fim:
                alvo[dia].append(etapa.rotulo_completo)

    return render_template(
        "provas/cronograma.html",
        inicio=inicio,
        fim=fim,
        dias=dias,
        semanas=semanas,
        opcoes_semanas=sorted(set(OPCOES_SEMANAS + [semanas])),
        meses_cabecalho=meses_cabecalho,
        linhas=linhas,
        df_dia=df_dia,
        sp_dia=sp_dia,
        dias_semana=DIAS_SEMANA_ABREV,
        anterior=inicio - timedelta(days=semanas * 7),
        proximo=inicio + timedelta(days=semanas * 7),
        semana_atual=_segunda_feira(hoje),
    )


@bp.route("/diario")
@requer("provas")
def diario():
    projeto_id = request.args.get("projeto", type=int)
    funcionario_id = request.args.get("funcionario", type=int)
    mes = request.args.get("mes")
    consulta = db.select(LancamentoDiario).join(EtapaProva).join(ProjetoProva).order_by(
        LancamentoDiario.data.desc(), ProjetoProva.ordem
    )
    if projeto_id:
        consulta = consulta.where(EtapaProva.projeto_id == projeto_id)
    if funcionario_id:
        consulta = consulta.where(LancamentoDiario.funcionario_id == funcionario_id)
    if mes:
        try:
            ano, numero = (int(p) for p in mes.split("-"))
            inicio_mes = date(ano, numero, 1)
            fim_mes = date(ano + 1, 1, 1) - timedelta(days=1) if numero == 12 else date(ano, numero + 1, 1) - timedelta(days=1)
            consulta = consulta.where(LancamentoDiario.data >= inicio_mes, LancamentoDiario.data <= fim_mes)
        except ValueError:
            mes = ""
    lancamentos = list(db.session.scalars(consulta.limit(500)))
    return render_template(
        "provas/diario.html",
        lancamentos=lancamentos,
        projetos=_projetos(apenas_ativos=False),
        funcionarios=list(db.session.scalars(db.select(Funcionario).order_by(Funcionario.nome))),
        filtro_projeto=projeto_id,
        filtro_funcionario=funcionario_id,
        filtro_mes=mes or "",
        tipos=TIPOS_LANCAMENTO,
    )


@bp.route("/diario/novo", methods=["GET", "POST"])
@bp.route("/diario/<int:lancamento_id>", methods=["GET", "POST"])
@requer("provas")
def diario_form(lancamento_id: int | None = None):
    lancamento = db.session.get(LancamentoDiario, lancamento_id) if lancamento_id else None
    if lancamento_id and not lancamento:
        abort(404)

    etapas = _etapas(apenas_ativos=False)
    funcionarios = list(db.session.scalars(db.select(Funcionario).filter_by(ativo=True).order_by(Funcionario.nome)))
    erros, avisos = [], []

    if request.method == "POST":
        etapa = db.session.get(EtapaProva, request.form.get("etapa_id", type=int) or 0)
        tipo = request.form.get("tipo") or "TRABALHO"
        so_uteis = request.form.get("so_uteis") == "1"
        tarefa = (request.form.get("tarefa") or "").strip() or None
        observacao = (request.form.get("observacao") or "").strip() or None
        realizada = request.form.get("realizada") == "1"
        escolhidos = _funcionarios_do_diario()
        inicio = _data_form("inicio") or _data_form("data")
        fim = _data_form("fim") or inicio

        if not etapa:
            erros.append("Selecione a etapa.")
        if not escolhidos:
            erros.append("Selecione pelo menos um diagramador.")
        if lancamento and len(escolhidos) != 1:
            erros.append("Na edição, escolha um diagramador.")
        if not inicio:
            erros.append("Informe a data inicial.")
        if inicio and fim and fim < inicio:
            erros.append("A data final não pode ser anterior à inicial.")
        if inicio and fim and (fim - inicio).days > 180:
            erros.append("O período não pode passar de 180 dias.")
        if tipo not in TIPOS_LANCAMENTO:
            erros.append("Tipo de lançamento inválido.")

        try:
            verba_informada = _decimal_form("verba")
        except ValueError as erro:
            verba_informada = None
            erros.append(str(erro))

        dias = _dias_do_periodo(inicio, fim, so_uteis) if inicio and fim and fim >= inicio else []
        if inicio and fim and not erros and not dias:
            erros.append("Não há dias nesse período com o filtro escolhido.")

        if etapa and dias and etapa.inicio and etapa.fim:
            fora = [d for d in dias if d < etapa.inicio or d > etapa.fim]
            if fora:
                avisos.append(
                    f"{len(fora)} dia(s) ficam fora do período da etapa "
                    f"({etapa.inicio.strftime('%d/%m/%Y')} a {etapa.fim.strftime('%d/%m/%Y')})."
                )
        if escolhidos and dias:
            ferias = ferias_no_intervalo(dias[0], dias[-1], apenas_ativos=True)
            ids = {f.id for f in escolhidos}
            for periodo in ferias:
                if periodo.funcionario_id in ids and any(periodo.inicio <= d <= periodo.fim for d in dias):
                    avisos.append(
                        f"{periodo.funcionario.nome} está de férias de "
                        f"{periodo.inicio.strftime('%d/%m/%Y')} a {periodo.fim.strftime('%d/%m/%Y')}."
                    )

        if not erros:
            if avisos and not request.form.get("confirmar"):
                erros.append("Há avisos para este lançamento. Confirme para salvar mesmo assim.")
            elif lancamento:
                funcionario = escolhidos[0]
                verba = verba_informada
                if verba is None:
                    verba = calcular_verba(etapa.local, tipo, funcionario.nome)
                lancamento.etapa = etapa
                lancamento.funcionario = funcionario
                lancamento.data = inicio
                lancamento.tipo = tipo
                lancamento.tarefa = tarefa
                lancamento.realizada = realizada
                lancamento.verba = verba
                lancamento.observacao = observacao
                lancamento.criado_por = current_user.login
                try:
                    registrar("Controle diário", f"Editar {funcionario.nome} em {etapa.rotulo_completo} ({inicio})")
                    db.session.commit()
                except IntegrityError:
                    db.session.rollback()
                    erros.append("Já existe lançamento deste funcionário nesta etapa neste dia.")
                else:
                    flash("Lançamento salvo.", "sucesso")
                    return redirect(url_for("provas.diario"))
            else:
                criados, ignorados = 0, 0
                for funcionario in escolhidos:
                    verba = calcular_verba(etapa.local, tipo, funcionario.nome)
                    if verba_informada is not None and len(escolhidos) == 1:
                        verba = verba_informada
                    for dia in dias:
                        ja = db.session.scalar(
                            db.select(LancamentoDiario).filter_by(
                                etapa_id=etapa.id, funcionario_id=funcionario.id, data=dia
                            )
                        )
                        if ja:
                            ignorados += 1
                            continue
                        db.session.add(
                            LancamentoDiario(
                                etapa=etapa,
                                funcionario=funcionario,
                                data=dia,
                                tipo=tipo,
                                tarefa=tarefa,
                                realizada=realizada,
                                verba=verba,
                                observacao=observacao,
                                criado_por=current_user.login,
                            )
                        )
                        criados += 1
                registrar(
                    "Controle diário",
                    f"{criados} lançamento(s) em {etapa.rotulo_completo} "
                    f"({inicio.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')})",
                )
                db.session.commit()
                if criados and ignorados:
                    flash(f"{criados} lançamento(s) salvo(s). {ignorados} já existiam e foram mantidos.", "sucesso")
                elif criados:
                    flash(f"{criados} lançamento(s) salvo(s).", "sucesso")
                else:
                    flash("Nada novo para salvar: esses dias já estavam lançados.", "aviso")
                return redirect(url_for("provas.diario"))

    return render_template(
        "provas/diario_form.html",
        lancamento=lancamento,
        etapas=etapas,
        funcionarios=funcionarios,
        tipos=TIPOS_LANCAMENTO,
        erros=erros,
        avisos=avisos,
        dados=request.form if request.method == "POST" else None,
        etapa_padrao=request.args.get("etapa", type=int),
        taxas_verba={chave: str(valor) for chave, valor in carregar_taxas().items()},
    )


@bp.route("/diario/<int:lancamento_id>/excluir", methods=["POST"])
@requer("provas")
def excluir_diario(lancamento_id: int):
    lancamento = db.session.get(LancamentoDiario, lancamento_id) or abort(404)
    registrar("Controle diário", f"Excluir {lancamento.funcionario.nome} em {lancamento.data}")
    db.session.delete(lancamento)
    db.session.commit()
    flash("Lançamento excluído.", "sucesso")
    return redirect(url_for("provas.diario"))


@bp.route("/relatorios")
@requer("provas")
def relatorios():
    projetos = _projetos(apenas_ativos=False)
    lancamentos = list(db.session.scalars(db.select(LancamentoDiario)))
    por_pessoa: dict[str, dict] = {}
    por_projeto: dict[int, dict] = {}
    for projeto in projetos:
        por_projeto[projeto.id] = {
            "projeto": projeto,
            "uteis": 0,
            "folgas": 0,
            "feriados": 0,
            "verba": Decimal("0"),
        }
    for item in lancamentos:
        nome = item.funcionario.nome
        pessoa = por_pessoa.setdefault(
            nome, {"nome": nome, "uteis": 0, "folgas": 0, "feriados": 0, "verba": Decimal("0")}
        )
        bloco = por_projeto.get(item.etapa.projeto_id)
        if item.tipo == "TRABALHO":
            pessoa["uteis"] += 1
            if bloco:
                bloco["uteis"] += 1
        elif item.tipo == "FERIADO_TRABALHADO":
            pessoa["feriados"] += 1
            if bloco:
                bloco["feriados"] += 1
        else:
            pessoa["folgas"] += 1
            if bloco:
                bloco["folgas"] += 1
        if item.verba:
            pessoa["verba"] += item.verba
            if bloco:
                bloco["verba"] += item.verba

    return render_template(
        "provas/relatorios.html",
        projetos=projetos,
        por_pessoa=sorted(por_pessoa.values(), key=lambda p: p["nome"]),
        por_projeto=list(por_projeto.values()),
        conflitos=conflitos_mesmo_local(_etapas(apenas_ativos=False)),
    )


@bp.route("/exportar")
@requer("provas")
def exportar():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="0F766E")

    def cabecalho(ws, larguras):
        for col, largura in enumerate(larguras, start=1):
            celula = ws.cell(row=1, column=col)
            celula.font, celula.fill = cabecalho_fonte, cabecalho_fundo
            celula.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = largura
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Projetos"
    ws.append(["Projeto", "Etapa", "Local", "Início", "Término", "Dias", "Equipe", "PCs adicionais", "Total PCs"])
    for projeto in _projetos(apenas_ativos=False):
        if not projeto.etapas:
            ws.append([projeto.nome, "", "", None, None, "", "", "", ""])
        for etapa in projeto.etapas:
            ws.append([
                projeto.nome,
                etapa.tipo,
                etapa.local or "",
                etapa.inicio,
                etapa.fim,
                etapa.dias or "",
                ", ".join(f.nome for f in etapa.equipe),
                etapa.pcs_adicionais,
                etapa.total_pcs,
            ])
            linha = ws.max_row
            ws.cell(row=linha, column=4).number_format = "DD/MM/YYYY"
            ws.cell(row=linha, column=5).number_format = "DD/MM/YYYY"
    cabecalho(ws, [28, 18, 10, 13, 13, 8, 36, 16, 12])

    diario_ws = wb.create_sheet("Controle diário")
    diario_ws.append(["Data", "Projeto", "Etapa", "Local", "Diagramador", "Tipo", "Tarefa", "Realizada", "Verba", "Observação"])
    for item in db.session.scalars(db.select(LancamentoDiario).order_by(LancamentoDiario.data)):
        diario_ws.append([
            item.data,
            item.etapa.projeto.nome,
            item.etapa.tipo,
            item.etapa.local or "",
            item.funcionario.nome,
            item.tipo_rotulo,
            item.tarefa or "",
            "Sim" if item.realizada else "Não",
            float(item.verba) if item.verba is not None else None,
            item.observacao or "",
        ])
        diario_ws.cell(row=diario_ws.max_row, column=1).number_format = "DD/MM/YYYY"
    cabecalho(diario_ws, [13, 24, 16, 10, 20, 18, 28, 12, 12, 36])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="controle_provas.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
