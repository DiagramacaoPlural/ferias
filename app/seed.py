"""Carga inicial do banco a partir da planilha original Férias_Final."""

import os
from datetime import date, datetime
from pathlib import Path

from .models import (
    PAPEL_ADMIN,
    PAPEL_GERENTE,
    PAPEL_SIMPLES,
    PAPEIS_FULL,
    STATUS_APROVADA,
    STATUS_PROGRAMADA,
    Ferias,
    Funcionario,
    Historico,
    Usuario,
    db,
)


def _primeiro_nome(nome: str) -> str:
    return (nome or "").strip().split()[0]


def _papel_do_nome(nome: str) -> str:
    chave = _primeiro_nome(nome).casefold()
    if chave in ("marcelo", "ricardo") or chave in PAPEIS_FULL:
        return PAPEL_GERENTE
    return PAPEL_SIMPLES


def _casar_funcionario_por_nome(nome: str, funcionarios: list[Funcionario]) -> Funcionario | None:
    chave = _primeiro_nome(nome).casefold()
    candidatos = [f for f in funcionarios if _primeiro_nome(f.nome).casefold() == chave]
    if len(candidatos) == 1:
        return candidatos[0]
    nome_cf = nome.casefold()
    for f in funcionarios:
        if f.nome.casefold() == nome_cf:
            return f
    return candidatos[0] if candidatos else None


def importar_usuarios_xlsx(caminho: Path) -> dict:
    """Cria usuários com o primeiro nome de cada funcionário da planilha Usuarios.xlsx."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, data_only=True, read_only=True)
    aba = wb[wb.sheetnames[0]]
    funcionarios = db.session.scalars(db.select(Funcionario)).all()
    resumo = {"usuarios": 0, "atualizados": 0, "sem_cadastro": []}
    senha_padrao = os.environ.get("USUARIO_SENHA_INICIAL") or "123"

    for i, linha in enumerate(aba.iter_rows(values_only=True), 1):
        if i == 1 or not linha or not linha[0]:
            continue
        nome = str(linha[0]).strip()
        email = str(linha[1]).strip() if len(linha) > 1 and linha[1] else ""
        login = _primeiro_nome(nome)
        if not login:
            continue
        func = _casar_funcionario_por_nome(nome, funcionarios)
        if not func:
            resumo["sem_cadastro"].append(nome)
            continue
        if email and "@" in email:
            func.email = email
        papel = _papel_do_nome(nome)
        full = papel != PAPEL_SIMPLES
        usuario = db.session.scalar(db.select(Usuario).filter(db.func.lower(Usuario.login) == login.lower()))
        if not usuario:
            usuario = Usuario(login=login, trocar_senha=True)
            usuario.definir_senha(senha_padrao)
            db.session.add(usuario)
            resumo["usuarios"] += 1
        else:
            resumo["atualizados"] += 1
        usuario.papel = papel
        if papel == PAPEL_SIMPLES:
            usuario.admin = False
        usuario.acessa_calendario = True
        usuario.acessa_cronograma = full
        usuario.funcionario_id = func.id
        usuario.ativo = True

    wb.close()
    db.session.add(
        Historico(
            usuario="sistema",
            acao="Usuários sincronizados",
            detalhe=f"{caminho.name}: {resumo['usuarios']} novos, {resumo['atualizados']} atualizados",
        )
    )
    db.session.commit()
    return resumo


def _sim(valor) -> bool:
    return str(valor or "").strip().upper() in ("SIM", "S", "TRUE", "1", "VERDADEIRO")


def _para_data(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str) and valor.strip():
        try:
            return datetime.strptime(valor.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def importar_planilha(caminho: Path) -> dict:
    """Importa funcionários, férias (aba CALENDARIO) e usuários (aba USUARIOS).

    Registros já existentes (mesmo nome/login) são ignorados, então a
    importação pode ser repetida sem duplicar dados.
    """
    from openpyxl import load_workbook

    wb = load_workbook(caminho, data_only=True, read_only=True)
    resumo = {"funcionarios": 0, "ferias": 0, "usuarios": 0}

    if "CALENDARIO" in wb.sheetnames:
        for linha in wb["CALENDARIO"].iter_rows(min_row=2, values_only=True):
            if not linha or not linha[0]:
                continue
            nome = str(linha[0]).strip()
            inicio = _para_data(linha[1] if len(linha) > 1 else None)
            dias = linha[3] if len(linha) > 3 else None
            aprovado = bool(linha[4]) if len(linha) > 4 else False

            funcionario = db.session.scalar(db.select(Funcionario).filter_by(nome=nome))
            if not funcionario:
                funcionario = Funcionario(nome=nome)
                db.session.add(funcionario)
                db.session.flush()
                resumo["funcionarios"] += 1

            if inicio and isinstance(dias, (int, float)) and dias > 0:
                ja_existe = db.session.scalar(
                    db.select(Ferias).filter_by(funcionario_id=funcionario.id, inicio=inicio)
                )
                if not ja_existe:
                    ferias = Ferias(
                        funcionario=funcionario,
                        status=STATUS_APROVADA if aprovado else STATUS_PROGRAMADA,
                        criado_por="importação",
                        aprovado_por="importação" if aprovado else None,
                        aprovado_em=datetime.now() if aprovado else None,
                    )
                    ferias.definir_periodo(inicio, int(dias))
                    db.session.add(ferias)
                    resumo["ferias"] += 1

    if "USUARIOS" in wb.sheetnames:
        for linha in wb["USUARIOS"].iter_rows(min_row=2, values_only=True):
            if not linha or not linha[0]:
                continue
            login = str(linha[0]).strip()
            if db.session.scalar(db.select(Usuario).filter(db.func.lower(Usuario.login) == login.lower())):
                continue
            usuario = Usuario(
                login=login,
                admin=_sim(linha[2] if len(linha) > 2 else None),
                acessa_calendario=_sim(linha[3] if len(linha) > 3 else None),
                acessa_cronograma=_sim(linha[4] if len(linha) > 4 else None),
                trocar_senha=True,
            )
            usuario.definir_senha(str(linha[1] if len(linha) > 1 and linha[1] is not None else "123"))
            db.session.add(usuario)
            resumo["usuarios"] += 1

    wb.close()
    db.session.add(
        Historico(
            usuario="sistema",
            acao="Importação",
            detalhe=f"{caminho.name}: {resumo['funcionarios']} funcionários, "
            f"{resumo['ferias']} férias, {resumo['usuarios']} usuários",
        )
    )
    db.session.commit()
    return resumo


def _planilhas_na_pasta(pasta: Path) -> list[Path]:
    arquivos = sorted(pasta.glob("*.xlsm")) + sorted(pasta.glob("*.xlsx"))
    return [p for p in arquivos if not p.name.startswith("~$")]


def _escolher_planilha(planilhas: list[Path], *palavras: str) -> Path | None:
    for planilha in planilhas:
        nome = planilha.name.lower()
        if any(palavra in nome for palavra in palavras):
            return planilha
    return None


_LINHA_PROGRAMACAO = (
    r"(?P<limite>\d{2}/\d{2}/\d{4})\s+20\s+10\s+Sim"
    r"(?P<saida>\d{2}/\d{2}/\d{4})"
    r".{0,20}?"
    r"(?P<aq_ini>\d{2}/\d{2}/\d{4})\s+a\s+(?P<aq_fim>\d{2}/\d{2}/\d{4})"
    r"(?:(?P<nome>[^\d\n]+?)(?P<matricula>\d{4,}))?"
)


def ler_programacao_pdf(caminho: Path) -> list[dict]:
    """Lê o PDF de Programação de Férias da folha (admissão = início do 1º aquisitivo)."""
    import re

    from pypdf import PdfReader

    texto = "\n".join((pagina.extract_text() or "") for pagina in PdfReader(str(caminho)).pages)
    pessoas, vistos = [], set()
    for m in re.finditer(_LINHA_PROGRAMACAO, texto):
        nome = " ".join((m.group("nome") or "").split())
        if not nome:
            continue
        chave = nome.casefold()
        if chave in vistos:
            continue
        vistos.add(chave)
        pessoas.append(
            {
                "nome": nome,
                "matricula": m.group("matricula"),
                "admissao": datetime.strptime(m.group("aq_ini"), "%d/%m/%Y").date(),
                "limite": datetime.strptime(m.group("limite"), "%d/%m/%Y").date(),
            }
        )
    return pessoas


def _casar_funcionario(nome_pdf: str, funcionarios: list[Funcionario]) -> Funcionario | None:
    primeiro = nome_pdf.split()[0].casefold()
    candidatos = [f for f in funcionarios if f.nome.split()[0].casefold() == primeiro]
    if len(candidatos) == 1:
        return candidatos[0]
    pdf = nome_pdf.casefold()
    for f in funcionarios:
        if f.nome.casefold() in pdf or pdf.startswith(f.nome.casefold()):
            return f
    return None


def importar_programacao_pdf(caminho: Path) -> dict:
    """Atualiza a data de admissão dos funcionários já cadastrados a partir do PDF da folha."""
    pessoas = ler_programacao_pdf(caminho)
    funcionarios = db.session.scalars(db.select(Funcionario)).all()
    resumo = {"lidos": len(pessoas), "atualizados": 0, "sem_cadastro": []}
    for pessoa in pessoas:
        func = _casar_funcionario(pessoa["nome"], funcionarios)
        if not func:
            resumo["sem_cadastro"].append(pessoa["nome"])
            continue
        if func.data_admissao != pessoa["admissao"]:
            func.data_admissao = pessoa["admissao"]
            resumo["atualizados"] += 1
    db.session.commit()
    return resumo


def nomes_abas(caminho: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True)
    nomes = list(wb.sheetnames)
    wb.close()
    return nomes


def garantir_usuarios_dos_funcionarios() -> dict:
    """Cria e liga um usuário para cada funcionário ativo (login = primeiro nome)."""
    senha_padrao = os.environ.get("USUARIO_SENHA_INICIAL") or "123"
    funcionarios = db.session.scalars(db.select(Funcionario).filter_by(ativo=True).order_by(Funcionario.nome)).all()
    resumo = {"criados": 0, "ligados": 0, "atualizados": 0}
    vistos: set[int] = set()
    for func in funcionarios:
        login = _primeiro_nome(func.nome)
        if not login:
            continue
        papel = _papel_do_nome(func.nome)
        full = papel != PAPEL_SIMPLES
        usuario = db.session.scalar(db.select(Usuario).filter(db.func.lower(Usuario.login) == login.lower()))
        if not usuario:
            usuario = Usuario(login=login, trocar_senha=True)
            usuario.definir_senha(senha_padrao)
            db.session.add(usuario)
            resumo["criados"] += 1
        else:
            resumo["atualizados"] += 1
        if usuario.funcionario_id != func.id:
            if usuario.funcionario_id is None:
                resumo["ligados"] += 1
            usuario.funcionario_id = func.id
        usuario.papel = papel
        if papel == PAPEL_SIMPLES:
            usuario.admin = False
        usuario.acessa_calendario = True
        usuario.acessa_cronograma = full or usuario.admin
        usuario.ativo = True
        vistos.add(usuario.id)

    for usuario in db.session.scalars(db.select(Usuario)):
        if usuario.id in vistos or usuario.funcionario_id:
            continue
        func = _casar_funcionario_por_nome(usuario.login, funcionarios)
        if not func:
            continue
        usuario.funcionario_id = func.id
        usuario.papel = _papel_do_nome(func.nome)
        resumo["ligados"] += 1

    if resumo["criados"] or resumo["ligados"]:
        db.session.add(
            Historico(
                usuario="sistema",
                acao="Usuários ligados aos funcionários",
                detalhe=f"{resumo['criados']} novos, {resumo['ligados']} ligados, {resumo['atualizados']} atualizados",
            )
        )
    db.session.commit()
    return resumo


def vincular_usuarios_aos_funcionarios() -> int:
    """Compatível com chamadas antigas: liga/cria usuários e devolve quantos foram ligados ou criados."""
    resumo = garantir_usuarios_dos_funcionarios()
    return resumo["criados"] + resumo["ligados"]


def inicializar_banco(pasta_planilha: Path | None) -> None:
    tem_usuario = db.session.scalar(db.select(db.func.count(Usuario.id)))
    if not tem_usuario:
        if pasta_planilha:
            planilhas = _planilhas_na_pasta(pasta_planilha)
            ferias = _escolher_planilha(planilhas, "feria") or (planilhas[0] if planilhas else None)
            if ferias:
                importar_planilha(ferias)

        if not db.session.scalar(db.select(db.func.count(Usuario.id))):
            admin = Usuario(login="admin", admin=True, papel=PAPEL_ADMIN, trocar_senha=True)
            admin.definir_senha(os.environ.get("ADMIN_SENHA_INICIAL") or "admin")
            db.session.add(admin)
            db.session.commit()

    if pasta_planilha:
        usuarios_xlsx = _escolher_planilha(_planilhas_na_pasta(pasta_planilha), "usuario")
        if usuarios_xlsx:
            importar_usuarios_xlsx(usuarios_xlsx)
    garantir_usuarios_dos_funcionarios()
